import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from hub.mcp_data import call

router = APIRouter(prefix="/api/yield", tags=["yield"])


class AnalyzeRequest(BaseModel):
    lot_ids: list[str]
    group_by: str
    yield_params: list[str]


@router.get("/lots")
async def list_lots():
    return {"lots": await call("ui_yield_lots")}


@router.get("/meta")
async def get_meta():
    meta = await call("ui_yield_meta")
    return {"grouping_columns": meta["grouping_columns"], "yield_params": meta["yield_params"]}


@router.post("/analyze")
async def analyze(req: AnalyzeRequest):
    return await call("ui_yield_analyze", **req.model_dump())


@router.post("/export")
async def export_excel(req: AnalyzeRequest):
    result = await call("ui_yield_analyze", **req.model_dump())

    wb = openpyxl.Workbook()
    _build_summary_sheet(wb.active, result, req)
    _build_raw_sheet(wb.create_sheet("Raw Data"), result, req)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=yield_analysis_{req.group_by}.xlsx"
        },
    )


# ── Excel helpers ──────────────────────────────────────────────────────────────

_H_FILL  = PatternFill("solid", fgColor="1E3A5F")
_H_FONT  = Font(bold=True, color="FFFFFF", size=10)
_SUB_FILL = PatternFill("solid", fgColor="2D4A6E")
_CENTER  = Alignment(horizontal="center", vertical="center")
_THIN    = Border(
    left=Side(style="thin", color="334155"),
    right=Side(style="thin", color="334155"),
    top=Side(style="thin", color="334155"),
    bottom=Side(style="thin", color="334155"),
)


def _hcell(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = _H_FILL
    c.font = _H_FONT
    c.alignment = _CENTER
    c.border = _THIN
    return c


def _dcell(ws, row, col, value, num_fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.alignment = _CENTER
    c.border = _THIN
    if num_fmt:
        c.number_format = num_fmt
    return c


def _build_summary_sheet(ws, result: dict, req: AnalyzeRequest):
    ws.title = "Summary"
    params = req.yield_params

    # Row 1: group headers
    _hcell(ws, 1, 1, "Group")
    _hcell(ws, 1, 2, "Count")
    col = 3
    for p in params:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 3)
        c = _hcell(ws, 1, col, p)
        col += 4

    # Row 2: sub-headers (Mean / Std / Min / Max)
    ws.cell(row=2, column=1).border = _THIN
    ws.cell(row=2, column=2).border = _THIN
    col = 3
    for _ in params:
        for sub in ["Mean", "Std", "Min", "Max"]:
            c = ws.cell(row=2, column=col, value=sub)
            c.fill = _SUB_FILL
            c.font = Font(bold=True, color="CCDDEE", size=9)
            c.alignment = _CENTER
            c.border = _THIN
            col += 1

    # Data rows
    for ridx, g in enumerate(result["groups"], 3):
        _dcell(ws, ridx, 1, g["group_value"])
        _dcell(ws, ridx, 2, g["count"])
        col = 3
        for p in params:
            s = g["stats"].get(p, {})
            _dcell(ws, ridx, col,     s.get("mean"), "0.00")
            _dcell(ws, ridx, col + 1, s.get("std"),  "0.00")
            _dcell(ws, ridx, col + 2, s.get("min"),  "0.00")
            _dcell(ws, ridx, col + 3, s.get("max"),  "0.00")
            col += 4

    # Column widths
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 8
    col_letter = "C"
    for i, _ in enumerate(params):
        for j in range(4):
            idx = 3 + i * 4 + j
            letter = openpyxl.utils.get_column_letter(idx)
            ws.column_dimensions[letter].width = 9
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 16

    # Freeze top 2 rows + first 2 cols
    ws.freeze_panes = "C3"


def _build_raw_sheet(ws, result: dict, req: AnalyzeRequest):
    ws.title = "Raw Data"
    headers = ["lot_id", "wafer_no", req.group_by] + req.yield_params
    for col, h in enumerate(headers, 1):
        _hcell(ws, 1, col, h)

    ridx = 2
    for g in result["groups"]:
        for w in g["wafers"]:
            _dcell(ws, ridx, 1, w.get("lot_id"))
            _dcell(ws, ridx, 2, w.get("wafer_no"))
            _dcell(ws, ridx, 3, g["group_value"])
            for colidx, p in enumerate(req.yield_params, 4):
                _dcell(ws, ridx, colidx, w.get(p), "0.00")
            ridx += 1

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    for i, _ in enumerate(req.yield_params):
        letter = openpyxl.utils.get_column_letter(4 + i)
        ws.column_dimensions[letter].width = 12
    ws.freeze_panes = "D2"
